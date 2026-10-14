"""
Invoice & Academic Marksheet Generator — Flask App
Run:  python app.py
Open: http://127.0.0.1:5000
"""

import os, json, re
from datetime import datetime, date, timedelta
from pathlib import Path

from flask import (Flask, render_template, request,
                   jsonify, send_file, abort)

# ── ReportLab ──────────────────────────────────────────────────────────────
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT, TA_CENTER, TA_LEFT
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                 Paragraph, Spacer, HRFlowable)

# ── openpyxl ───────────────────────────────────────────────────────────────
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────────────────────────────────
app = Flask(__name__)
OUTPUT_DIR = Path("outputs")
OUTPUT_DIR.mkdir(exist_ok=True)

# ── Brand colours ──────────────────────────────────────────────────────────
BRAND  = colors.HexColor("#00A7C8")
DARK   = colors.HexColor("#071827")
LIGHT  = colors.HexColor("#eafaff")
GREY   = colors.HexColor("#666666")
RED_   = colors.HexColor("#d64545")
GREEN_ = colors.HexColor("#2e9e44")
WHITE  = colors.white
BLACK  = colors.black

_styles = getSampleStyleSheet()

def PS(name="N", size=9, bold=False, color=BLACK, align=TA_LEFT):
    return ParagraphStyle(
        name, parent=_styles["Normal"],
        fontSize=size, leading=size * 1.35,
        fontName="Helvetica-Bold" if bold else "Helvetica",
        textColor=color, alignment=align,
    )

def P(text, **kw):
    return Paragraph(str(text), PS(**kw))

def safe(s):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(s))[:50]

# ══════════════════════════════════════════════════════════════════════════
#  HELPERS
# ══════════════════════════════════════════════════════════════════════════
def num_to_words(n):
    ones = ["","One","Two","Three","Four","Five","Six","Seven","Eight","Nine",
            "Ten","Eleven","Twelve","Thirteen","Fourteen","Fifteen","Sixteen",
            "Seventeen","Eighteen","Nineteen"]
    tens = ["","","Twenty","Thirty","Forty","Fifty","Sixty","Seventy","Eighty","Ninety"]
    def h(n):
        if n < 20:    return ones[n]
        if n < 100:   return tens[n//10] + (" " + ones[n%10] if n%10 else "")
        if n < 1000:  return ones[n//100] + " Hundred" + (" and " + h(n%100) if n%100 else "")
        if n < 1e5:   return h(n//1000) + " Thousand" + (" " + h(n%1000) if n%1000 else "")
        if n < 1e7:   return h(int(n//1e5)) + " Lakh" + (" " + h(int(n%1e5)) if int(n%1e5) else "")
        return h(int(n//1e7)) + " Crore" + (" " + h(int(n%1e7)) if int(n%1e7) else "")
    r = int(n); p = round((n - r) * 100)
    return h(r) + " Rupees" + (" and " + h(p) + " Paise" if p else "") + " Only"

GRADE_SCALE = [
    (90,"O",10),(80,"A+",9),(70,"A",8),(60,"B+",7),(50,"B",6),(40,"C",5),(0,"F",0)
]

def get_grade(pct):
    for threshold, letter, gp in GRADE_SCALE:
        if pct >= threshold:
            return letter, gp
    return "F", 0

# ══════════════════════════════════════════════════════════════════════════
#  INVOICE PDF
# ══════════════════════════════════════════════════════════════════════════
def build_invoice_pdf(data: dict) -> Path:
    company      = data["company"]
    client       = data["client"]
    meta         = data["meta"]
    line_items   = data["line_items"]
    tax_mode     = data.get("tax_mode", "cgst_sgst")
    sym          = meta.get("currency_symbol", "Rs.")
    invoice_date = datetime.strptime(meta["invoice_date"], "%Y-%m-%d").date()
    due_date     = invoice_date + timedelta(days=int(meta.get("due_days", 30)))

    # ── Calculations ──
    grand_subtotal = grand_discount = grand_tax = grand_total = 0
    computed = []
    for item in line_items:
        gross    = float(item["qty"]) * float(item["rate"])
        disc     = round(gross * float(item.get("discount_pct", 0)) / 100, 2)
        taxable  = round(gross - disc, 2)
        tax_amt  = round(taxable * float(item.get("tax_pct", 18)) / 100, 2)
        lt       = round(taxable + tax_amt, 2)
        grand_subtotal += gross; grand_discount += disc
        grand_tax += tax_amt;   grand_total += lt
        computed.append({**item, "gross": gross, "disc_amt": disc,
                         "taxable": taxable, "tax_amt": tax_amt, "line_total": lt})

    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = OUTPUT_DIR / f"Invoice_{safe(meta.get('invoice_no','INV'))}_{ts}.pdf"

    doc = SimpleDocTemplate(str(path), pagesize=A4,
                             leftMargin=15*mm, rightMargin=15*mm,
                             topMargin=15*mm, bottomMargin=20*mm)
    story = []

    # Header
    hdr = Table([
        [P(company["name"], size=16, bold=True, color=BRAND),
         P("<b>INVOICE</b>", size=20, bold=True, color=DARK, align=TA_RIGHT)],
        [P(company.get("tagline",""), size=8, color=GREY),
         P(f"No: <b>{meta.get('invoice_no','')}</b>", size=9, align=TA_RIGHT)],
        [P(company.get("address",""), size=8, color=GREY),
         P(f"Date: {invoice_date.strftime('%d %b %Y')}", size=9, align=TA_RIGHT)],
        [P(f"GSTIN: {company.get('gstin','')}  |  {company.get('email','')}", size=8, color=GREY),
         P(f"Due: <b>{due_date.strftime('%d %b %Y')}</b>", size=9, color=RED_, align=TA_RIGHT)],
    ], colWidths=[95*mm, 85*mm])
    hdr.setStyle(TableStyle([
        ("VALIGN", (0,0),(-1,-1),"TOP"),
        ("LINEBELOW",(0,-1),(-1,-1),1.5,BRAND),
    ]))
    story += [hdr, Spacer(1, 5*mm)]

    # Bill To
    bill = Table([
        [P("BILL TO", size=8, bold=True, color=BRAND), P("", size=8)],
        [P(f"<b>{client['name']}</b>", size=10),
         P(f"<b>Contact:</b> {client.get('contact','')}", size=9)],
        [P(client.get("address",""), size=9, color=GREY),
         P(f"<b>GSTIN:</b> {client.get('gstin','')}", size=9, color=GREY)],
        [P(f"Email: {client.get('email','')}", size=8, color=GREY),
         P(f"Place of Supply: {meta.get('place_of_supply','')}", size=8, color=GREY)],
    ], colWidths=[95*mm, 85*mm])
    bill.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),LIGHT),
        ("BOX",(0,0),(-1,-1),0.5,colors.HexColor("#c0d8f0")),
        ("INNERGRID",(0,0),(-1,-1),0.3,colors.HexColor("#dde8f5")),
        ("LEFTPADDING",(0,0),(-1,-1),5),
        ("TOPPADDING",(0,0),(-1,-1),3),
        ("BOTTOMPADDING",(0,0),(-1,-1),3),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
    ]))
    story += [bill, Spacer(1, 5*mm)]

    # Line items table — sums to exactly 180mm
    if tax_mode == "cgst_sgst":
        heads  = ["#","Description","Qty","Rate","Discount","Taxable","CGST(9%)","SGST(9%)","Total"]
        col_w  = [6, 54, 10, 18, 20, 18, 14, 14, 26]
    else:
        heads  = ["#","Description","Qty","Rate","Discount","Taxable","IGST","Total"]
        col_w  = [6, 62, 10, 20, 20, 18, 18, 26]

    rows = [[P(h, size=8, bold=True, color=WHITE, align=TA_CENTER) for h in heads]]
    for idx, item in enumerate(computed, 1):
        ht = item["tax_amt"] / 2
        base = [
            P(str(idx), size=8, align=TA_CENTER),
            P(item["description"], size=8),
            P(str(item["qty"]), size=8, align=TA_CENTER),
            P(f"{sym}{float(item['rate']):,.0f}", size=8, align=TA_RIGHT),
            P(f"{item.get('discount_pct',0)}% {sym}{item['disc_amt']:,.0f}", size=8, align=TA_RIGHT),
            P(f"{sym}{item['taxable']:,.0f}", size=8, align=TA_RIGHT),
        ]
        if tax_mode == "cgst_sgst":
            base += [P(f"{sym}{ht:,.0f}", size=8, align=TA_RIGHT),
                     P(f"{sym}{ht:,.0f}", size=8, align=TA_RIGHT)]
        else:
            base.append(P(f"{item.get('tax_pct',18)}% {sym}{item['tax_amt']:,.0f}", size=8, align=TA_RIGHT))
        base.append(P(f"{sym}{item['line_total']:,.0f}", size=9, bold=True, align=TA_RIGHT))
        rows.append(base)

    items_tbl = Table(rows, colWidths=[w*mm for w in col_w], repeatRows=1)
    items_tbl.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),DARK), ("TEXTCOLOR",(0,0),(-1,0),WHITE),
        ("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#c0d8f0")),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE,LIGHT]),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4),
    ]))
    story += [items_tbl, Spacer(1, 4*mm)]

    # Totals
    tax_label = "CGST + SGST" if tax_mode == "cgst_sgst" else f"IGST {line_items[0].get('tax_pct',18)}%"
    tot_rows = [
        ["Gross Amount",    f"{sym} {grand_subtotal:>12,.2f}"],
        ["Total Discount",  f"- {sym} {grand_discount:>11,.2f}"],
        ["Taxable Amount",  f"{sym} {(grand_subtotal-grand_discount):>12,.2f}"],
        [tax_label,         f"{sym} {grand_tax:>12,.2f}"],
        [P("<b>GRAND TOTAL</b>", size=11, bold=True, color=WHITE, align=TA_RIGHT),
         P(f"<b>{sym} {grand_total:>12,.2f}</b>", size=11, bold=True, color=WHITE, align=TA_RIGHT)],
    ]
    tot_tbl = Table([["", r[0], r[1]] for r in tot_rows],
                    colWidths=[90*mm, 58*mm, 32*mm])
    tot_tbl.setStyle(TableStyle([
        ("ALIGN",(1,0),(-1,-1),"RIGHT"),
        ("GRID",(1,0),(-1,-2),0.3,colors.HexColor("#c0d8f0")),
        ("BACKGROUND",(0,-1),(-1,-1),DARK),
        ("LINEABOVE",(1,-1),(-1,-1),1.5,BRAND),
        ("TOPPADDING",(0,0),(-1,-1),3), ("BOTTOMPADDING",(0,0),(-1,-1),3),
    ]))
    story += [tot_tbl, Spacer(1, 3*mm)]
    story.append(P(f"<b>Amount in Words:</b> <i>{num_to_words(grand_total)}</i>", size=9))
    story.append(Spacer(1, 5*mm))

    # Notes + Bank
    nb_tbl = Table([[
        [P("Notes", size=8, bold=True, color=BRAND),
         P(meta.get("notes",""), size=8, color=GREY)],
        [P("Bank Details", size=8, bold=True, color=BRAND),
         P(company.get("bank_details",""), size=8, color=GREY)],
    ]], colWidths=[90*mm, 90*mm])
    nb_tbl.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("BOX",(0,0),(0,0),0.5,colors.HexColor("#c0d8f0")),
        ("BOX",(1,0),(1,0),0.5,colors.HexColor("#c0d8f0")),
        ("BACKGROUND",(0,0),(0,0),LIGHT), ("BACKGROUND",(1,0),(1,0),LIGHT),
        ("LEFTPADDING",(0,0),(-1,-1),5),
        ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4),
    ]))
    story += [nb_tbl, Spacer(1, 6*mm)]
    story.append(HRFlowable(width="100%", thickness=1, color=BRAND))
    story.append(Spacer(1, 2*mm))
    story.append(P(
        f"{company.get('website','')}  |  {company.get('email','')}  |  {company.get('phone','')}",
        size=8, color=GREY, align=TA_CENTER))
    story.append(P("Computer-generated invoice — does not require a physical signature.",
                   size=7, color=GREY, align=TA_CENTER))

    doc.build(story)
    return path


# ══════════════════════════════════════════════════════════════════════════
#  MARKSHEET PDF  (one student)
# ══════════════════════════════════════════════════════════════════════════
def build_marksheet_pdf(institution, subjects, student, all_students) -> Path:
    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = OUTPUT_DIR / f"Marksheet_{safe(student['roll_no'])}_{safe(student['name'])}_{ts}.pdf"

    doc = SimpleDocTemplate(str(path), pagesize=A4,
                             leftMargin=15*mm, rightMargin=15*mm,
                             topMargin=12*mm, bottomMargin=15*mm)
    story = []

    # Header
    hdr = Table([[
        P(institution.get("logo_text","INST"), size=26, bold=True, color=BRAND),
        [P(institution["name"], size=13, bold=True, color=DARK),
         P(institution.get("subtitle",""), size=9, color=GREY),
         P(institution.get("address",""), size=8, color=GREY),
         P(f"Affiliation: {institution.get('affiliation','')}", size=7, color=GREY)],
    ]], colWidths=[22*mm, 158*mm])
    hdr.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LINEBELOW",(0,-1),(-1,-1),2,BRAND),
    ]))
    story += [hdr, Spacer(1, 3*mm)]

    banner = Table([[P(institution.get("exam_name","Examination"), size=10,
                       bold=True, color=WHITE, align=TA_CENTER)]],
                   colWidths=[180*mm])
    banner.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),DARK),
        ("TOPPADDING",(0,0),(-1,-1),5), ("BOTTOMPADDING",(0,0),(-1,-1),5),
    ]))
    story += [banner, Spacer(1, 4*mm)]

    det = Table([
        [P("Roll No:", size=9, bold=True),  P(student["roll_no"], size=9),
         P("Programme:", size=9, bold=True), P(institution.get("program",""), size=9)],
        [P("Name:", size=9, bold=True),      P(student["name"], size=9, bold=True),
         P("Semester:", size=9, bold=True),  P(institution.get("semester",""), size=9)],
        [P("Father:", size=9, bold=True),    P(student.get("father",""), size=9),
         P("Date of Birth:", size=9, bold=True), P(student.get("dob",""), size=9)],
    ], colWidths=[28*mm, 57*mm, 33*mm, 62*mm])
    det.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),LIGHT),
        ("BOX",(0,0),(-1,-1),0.5,colors.HexColor("#c0d8f0")),
        ("INNERGRID",(0,0),(-1,-1),0.3,colors.HexColor("#dde8f5")),
        ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING",(0,0),(-1,-1),5),
    ]))
    story += [det, Spacer(1, 5*mm)]

    # Marks table — col widths sum = 169mm < 180mm ✓
    col_w_ms = [16, 52, 12, 18, 18, 13, 10, 14, 16]
    m_heads  = ["Code","Subject Name","Max","Obtained","Percent","Grade","GP","Cr","Status"]
    m_rows   = [[P(h, size=8, bold=True, color=WHITE, align=TA_CENTER) for h in m_heads]]

    for s in student["subjects"]:
        m_rows.append([
            P(s["code"], size=8, align=TA_CENTER),
            P(s["name"], size=8),
            P(str(s["max"]), size=8, align=TA_CENTER),
            P(str(s["obtained"]), size=9, bold=True, align=TA_CENTER),
            P(f"{s['pct']}%", size=8, align=TA_CENTER),
            P(s["grade"], size=9, bold=True, align=TA_CENTER,
              color=BRAND if s["grade"] not in ("F","C") else RED_),
            P(str(s["gp"]), size=8, align=TA_CENTER),
            P(str(s["credits"]), size=8, align=TA_CENTER),
            P("PASS" if s["passed"] else "FAIL", size=8, bold=True,
              color=GREEN_ if s["passed"] else RED_, align=TA_CENTER),
        ])

    m_tbl = Table(m_rows, colWidths=[w*mm for w in col_w_ms], repeatRows=1)
    m_tbl.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),DARK), ("TEXTCOLOR",(0,0),(-1,0),WHITE),
        ("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#c0d8f0")),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE,LIGHT]),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4),
    ]))
    story += [m_tbl, Spacer(1, 5*mm)]

    # Result summary bar — col widths sum = 180mm ✓
    res_c = GREEN_ if student["result"] == "PASS" else RED_
    summary = Table([[
        P(f"Total: {student['total_obtained']}/{student['total_max']}", size=10, bold=True, color=WHITE),
        P(f"Percentage: {student['overall_pct']}%", size=10, bold=True, color=WHITE),
        P(f"CGPA: {student['cgpa']}", size=10, bold=True, color=WHITE),
        P(f"Grade: {student['overall_grade']}", size=10, bold=True, color=WHITE),
        P(f"Rank: #{student['rank']}", size=10, bold=True, color=WHITE),
        P(student["result"], size=11, bold=True, color=WHITE, align=TA_CENTER),
    ]], colWidths=[35*mm, 38*mm, 28*mm, 24*mm, 23*mm, 32*mm])
    summary.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-2,-1),DARK),
        ("BACKGROUND",(-1,0),(-1,-1),res_c),
        ("TOPPADDING",(0,0),(-1,-1),6), ("BOTTOMPADDING",(0,0),(-1,-1),6),
        ("LEFTPADDING",(0,0),(-1,-1),4), ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
    ]))
    story.append(summary)

    if student.get("ai_remark"):
        story += [Spacer(1,4*mm),
                  P(f"<b>Remarks:</b> <i>{student['ai_remark']}</i>", size=9, color=DARK)]

    story.append(Spacer(1, 10*mm))
    sig = Table([[
        P("_______________________", size=9, align=TA_CENTER),
        P("_______________________", size=9, align=TA_CENTER),
        P("_______________________", size=9, align=TA_CENTER),
    ],[
        P("Class Teacher", size=8, color=GREY, align=TA_CENTER),
        P("Controller of Exams", size=8, color=GREY, align=TA_CENTER),
        P("Principal / Director", size=8, color=GREY, align=TA_CENTER),
    ]], colWidths=[60*mm, 60*mm, 60*mm])
    story += [sig, Spacer(1, 3*mm)]
    story.append(HRFlowable(width="100%", thickness=0.5, color=BRAND))
    story.append(P("Computer-generated — any alteration renders this marksheet invalid.",
                   size=7, color=GREY, align=TA_CENTER))
    doc.build(story)
    return path


# ══════════════════════════════════════════════════════════════════════════
#  CLASS RESULT EXCEL
# ══════════════════════════════════════════════════════════════════════════
def build_class_excel(institution, subjects, students) -> Path:
    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = OUTPUT_DIR / f"ClassResult_{safe(institution['name'])}_{ts}.xlsx"

    wb = Workbook(); ws = wb.active
    assert ws is not None
    ws.title = "Class Result"

    def xf(h): return PatternFill("solid", fgColor=h.lstrip("#"))
    def xfont(bold=False, color="000000", size=10):
        return Font(bold=bold, color=color, size=size, name="Calibri")
    thin = Side(style="thin", color="C0D8F0")
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lft  = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    last_col = get_column_letter(3 + len(subjects) + 7)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"{institution['name']} | {institution.get('exam_name','')}"
    ws["A1"].font = xfont(bold=True, color="1a2940", size=13)
    ws["A1"].alignment = ctr; ws["A1"].fill = xf("f0f7ff")
    ws.row_dimensions[1].height = 22

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = f"{institution.get('program','')} | {institution.get('semester','')} | Students: {len(students)}"
    ws["A2"].font = xfont(color="3BA4E8", size=10)
    ws["A2"].alignment = ctr
    ws.append([])

    heads = (["Roll No","Name","Father"] +
             [f"{s['code']}\n(/{s['max_marks']})" for s in subjects] +
             ["Total","Max","Percentage","CGPA","Grade","Rank","Result"])
    ws.append(heads)
    hr = ws.max_row
    for ci, _ in enumerate(heads, 1):
        c = ws.cell(hr, ci)
        c.font = xfont(bold=True, color="FFFFFF", size=9)
        c.fill = xf("1a2940"); c.alignment = ctr; c.border = bdr
    ws.row_dimensions[hr].height = 28

    for ii, stu in enumerate(sorted(students, key=lambda x: x["rank"])):
        row = ([stu["roll_no"], stu["name"], stu.get("father","")] +
               [s["obtained"] for s in stu["subjects"]] +
               [stu["total_obtained"], stu["total_max"], stu["overall_pct"],
                stu["cgpa"], stu["overall_grade"], stu["rank"], stu["result"]])
        ws.append(row)
        r  = ws.max_row
        rf = xf("F0F7FF") if ii % 2 == 0 else xf("FFFFFF")
        for ci in range(1, len(row)+1):
            c = ws.cell(r, ci)
            c.border = bdr; c.alignment = ctr; c.fill = rf; c.font = xfont(size=9)
            si = 4; se = 4 + len(subjects)
            if si <= ci < se:
                s = stu["subjects"][ci - si]
                if not s["passed"]:
                    c.fill = xf("FFE4E4")
                    c.font = Font(bold=True, color="d64545", size=9, name="Calibri")
            if ci == len(row):
                is_pass = stu["result"] == "PASS"
                c.font  = Font(bold=True, color="2e9e44" if is_pass else "d64545",
                               size=9, name="Calibri")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    for i in range(len(subjects)):
        ws.column_dimensions[get_column_letter(4+i)].width = 9
    for j, w in enumerate([8,8,12,8,8,6,16], len(subjects)+4):
        ws.column_dimensions[get_column_letter(j)].width = w

    # Statistics sheet
    ws2 = wb.create_sheet("Statistics")
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 35
    ws2.append(["Metric","Value"])
    for ci in [1,2]:
        c = ws2.cell(1,ci)
        c.font = xfont(bold=True,color="FFFFFF"); c.fill=xf("1a2940")
        c.alignment=ctr; c.border=bdr
    n      = len(students)
    passed = sum(1 for s in students if s["result"]=="PASS")
    avp    = round(sum(s["overall_pct"] for s in students)/n, 2) if n else 0
    avc    = round(sum(s["cgpa"] for s in students)/n, 2) if n else 0
    topper = min(students, key=lambda x: x["rank"]) if students else {}
    stats  = [
        ("Total Students", n), ("Passed", passed), ("Failed", n-passed),
        ("Pass %", f"{round(passed/n*100,1)}%" if n else "N/A"),
        ("Class Average %", avp), ("Class Average CGPA", avc),
        ("Topper", topper.get("name","")),
        ("Topper Score", f"{topper.get('overall_pct','')}% | CGPA {topper.get('cgpa','')}"),
    ]
    for k, v in stats:
        ws2.append([k, str(v)]); r2 = ws2.max_row
        for ci in [1,2]:
            c=ws2.cell(r2,ci); c.border=bdr; c.alignment=lft; c.font=xfont(size=10)

    wb.save(str(path))
    return path


# ══════════════════════════════════════════════════════════════════════════
#  MARKSHEET CALCULATIONS HELPER
# ══════════════════════════════════════════════════════════════════════════
def compute_marksheets(subjects, students_raw):
    total_credits = sum(s["credits"] for s in subjects)
    computed = []
    for stu in students_raw:
        subj_det, failed = [], []
        tot_obt = tot_max = wgp = 0
        for i, subj in enumerate(subjects):
            obt   = int(stu["marks"][i])
            maxm  = subj["max_marks"]
            pct   = round(obt / maxm * 100, 1)
            grade, gp = get_grade(pct)
            passed = obt >= subj["pass_marks"] and grade != "F"
            if not passed: failed.append(subj["code"])
            wgp += gp * subj["credits"]
            tot_obt += obt; tot_max += maxm
            subj_det.append({
                "code": subj["code"], "name": subj["name"],
                "max": maxm, "obtained": obt, "pct": pct,
                "grade": grade, "gp": gp,
                "credits": subj["credits"], "passed": passed,
            })
        ovr_pct   = round(tot_obt / tot_max * 100, 2)
        cgpa      = round(wgp / total_credits, 2)
        ovr_grade, _ = get_grade(ovr_pct)
        result    = "PASS" if not failed else f"FAIL ({', '.join(failed)})"
        computed.append({
            "roll_no": stu["roll_no"], "name": stu["name"],
            "father": stu.get("father",""), "dob": stu.get("dob",""),
            "subjects": subj_det, "total_obtained": tot_obt, "total_max": tot_max,
            "overall_pct": ovr_pct, "cgpa": cgpa,
            "overall_grade": ovr_grade, "result": result,
            "failed_subjects": failed, "ai_remark": "",
        })
    computed.sort(key=lambda x: -x["overall_pct"])
    for rank, s in enumerate(computed, 1): s["rank"] = rank
    return computed


def build_page_report():
    pages = [
        {
            "page_name": "home.html",
            "route": "/",
            "label": "Landing Page",
            "summary": "Introduces Anshita's Document Studio, explains the core value, and sends users into the workspace.",
            "items": [
                {
                    "element": "Hero headline",
                    "text": "Anshita's Document Studio.",
                    "notes": "Main product identity shown in the top hero section.",
                },
                {
                    "element": "Hero description",
                    "text": "A polished document automation workspace for invoices, academic marksheets, PDF exports, Excel summaries, and clean project reporting.",
                    "notes": "Short explanation of the product goal.",
                },
                {
                    "element": "Primary CTA",
                    "text": "Open Workspace",
                    "notes": "Takes the user to the interactive builder page.",
                },
                {
                    "element": "Secondary CTA",
                    "text": "Explore Features",
                    "notes": "Scrolls to the feature section on the same page.",
                },
                {
                    "element": "Feature section heading",
                    "text": "Advanced document automation, ready for demos.",
                    "notes": "Introduces the three key feature cards.",
                },
                {
                    "element": "Mentorship heading",
                    "text": "Built for Anshita's workflow",
                    "notes": "Highlights the personalized project identity and workflow.",
                },
                {
                    "element": "Bottom CTA",
                    "text": "Go to Workspace",
                    "notes": "Repeat action link at the end of the page.",
                },
            ],
        },
        {
            "page_name": "index.html",
            "route": "/workspace",
            "label": "Workspace Page",
            "summary": "Hosts the real invoice generator and marksheet generator flows, including backend PDF and Excel generation.",
            "items": [
                {
                    "element": "Module tab 1",
                    "text": "Invoices",
                    "notes": "Switches the workspace into invoice mode.",
                },
                {
                    "element": "Module tab 2",
                    "text": "Marksheets",
                    "notes": "Switches the workspace into marksheet mode.",
                },
                {
                    "element": "Invoice module title",
                    "text": "Invoice Generator",
                    "notes": "Label shown above the invoice form.",
                },
                {
                    "element": "Invoice module subtitle",
                    "text": "Create branded, tax-aware invoice PDFs with live totals.",
                    "notes": "Explains the invoice workflow in one line.",
                },
                {
                    "element": "Invoice mode controls",
                    "text": "Manual Entry / Bulk Upload",
                    "notes": "Lets the user choose how invoice data is provided.",
                },
                {
                    "element": "Invoice action button",
                    "text": "Generate Invoice PDF",
                    "notes": "Triggers invoice PDF generation.",
                },
                {
                    "element": "Invoice success heading",
                    "text": "Generation Successful",
                    "notes": "Shown after invoice files are created.",
                },
                {
                    "element": "Marksheet module title",
                    "text": "Marksheet Generator",
                    "notes": "Label shown above the marksheet form.",
                },
                {
                    "element": "Marksheet module subtitle",
                    "text": "Generate ranked marksheets and Excel class summaries.",
                    "notes": "Explains the marksheet workflow in one line.",
                },
                {
                    "element": "Marksheet action button",
                    "text": "Generate Reports",
                    "notes": "Triggers marksheet and class summary generation.",
                },
                {
                    "element": "Preview heading",
                    "text": "Class Performance Preview",
                    "notes": "Shows computed preview rows before downloads.",
                },
                {
                    "element": "Result heading",
                    "text": "Reports Generated",
                    "notes": "Shown after class report files are created.",
                },
            ],
        },
    ]

    return {
        "title": "Page Copy Report",
        "generated_at": datetime.now().strftime("%B %d, %Y at %I:%M %p"),
        "page_count": len(pages),
        "item_count": sum(len(page["items"]) for page in pages),
        "pages": pages,
    }


# ══════════════════════════════════════════════════════════════════════════
#  ROUTES
# ══════════════════════════════════════════════════════════════════════════
@app.route("/")
def home():
    return render_template("home.html")


@app.route("/home")
def home_page():
    return render_template("home.html")


@app.route("/workspace")
def index():
    return render_template("index.html")


@app.route("/report")
def page_report():
    return render_template("report.html", report=build_page_report())


@app.route("/api/generate-invoice", methods=["POST"])
def generate_invoice():
    try:
        data = request.get_json(force=True)
        path = build_invoice_pdf(data)
        return jsonify({"status": "ok", "filename": path.name, "download": f"/download/{path.name}"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400


@app.route("/api/generate-marksheet", methods=["POST"])
def generate_marksheet():
    try:
        data        = request.get_json(force=True)
        institution = data["institution"]
        subjects    = data["subjects"]
        students_raw= data["students"]
        mode        = data.get("mode", "all")   # "single" | "all" | "excel"

        students = compute_marksheets(subjects, students_raw)

        files = []
        if mode in ("single", "all"):
            target = students if mode == "all" else [students[0]]
            for stu in target:
                p = build_marksheet_pdf(institution, subjects, stu, students)
                files.append({"filename": p.name, "download": f"/download/{p.name}",
                               "student": stu["name"]})

        if mode in ("excel", "all"):
            p = build_class_excel(institution, subjects, students)
            files.append({"filename": p.name, "download": f"/download/{p.name}",
                          "student": "Class Summary (Excel)"})

        # Return computed results for preview table
        preview = [{
            "roll_no": s["roll_no"], "name": s["name"],
            "total": f"{s['total_obtained']}/{s['total_max']}",
            "pct": s["overall_pct"], "cgpa": s["cgpa"],
            "grade": s["overall_grade"], "rank": s["rank"], "result": s["result"],
        } for s in students]

        return jsonify({"status": "ok", "files": files, "preview": preview})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400


@app.route("/download/<filename>")
def download(filename):
    path = OUTPUT_DIR / filename
    if not path.exists() or not path.is_file():
        abort(404)
    return send_file(str(path.resolve()), as_attachment=True, download_name=filename)


if __name__ == "__main__":
    app.run(debug=True, port=5000)
